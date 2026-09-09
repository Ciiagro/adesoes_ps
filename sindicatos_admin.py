"""
Lógica do módulo "Sindicatos Rurais x Municípios": listagem/filtro, edição
com histórico de auditoria, e geração da planilha para os setores.
As rotas Flask (em app.py) só chamam essas funções.
"""
import re
from datetime import datetime, date, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
)
from sqlalchemy.orm import joinedload
from sqlalchemy import or_, func

from models import SindicatoRural, MunicipioSindicato, HistoricoSindicato, ProcessoEleitoral, DiretoriaMembro
import supabase_storage

# Nota: nos joinedload(MunicipioSindicato.sindicato) abaixo, encadeamos
# .defer(SindicatoRural.foto_presidente).defer(...tipo) pra NUNCA puxar o
# bytea antigo da foto em telas de lista/mapa — era isso que inflava o
# egress. A foto atual mora no Storage (foto_presidente_url), que é leve
# (é só uma string).

# Campos editáveis do cadastro de um município (nome do atributo -> rótulo
# amigável, usado tanto no formulário quanto no histórico/auditoria).
CAMPOS_EDITAVEIS = {
    # --- Sindicato: nome dos cargos + contato do sindicato ---
    "presidente_nome": "Presidente do Sindicato",
    "secretario_nome": "Secretário(a) do Sindicato",
    "telefone1": "Telefone 1 (Sindicato)",
    "telefone2": "Telefone 2 (Sindicato)",
    "email1": "E-mail 1 (Sindicato)",
    "email2": "E-mail 2 (Sindicato)",
    "email3": "E-mail 3 (Sindicato)",
    "email4": "E-mail 4 (Sindicato)",
    # --- Vice-Presidente Regional: pessoa distinta (às vezes também é
    # presidente de algum sindicato específico), com contato próprio ---
    "vice_presidente_texto": "Vice-Presidente Regional",
    "telefone_vice_presidente": "Telefone do Vice-Presidente Regional",
    "email_vice_presidente": "E-mail do Vice-Presidente Regional",
}

COLUNAS_EXPORTACAO = [
    ("cod_ibge", "Cod_IBGE"),
    ("nome", "Município"),
    ("regiao_faec", "Região FAEC"),
    ("regiao_sebrae", "Região Sebrae"),
    ("regiao_planejamento", "Região Planejamento"),
    ("_sindicato_nome", "Sindicato Responsável"),
    ("presidente_nome", "Presidente do Sindicato"),
    ("secretario_nome", "Secretário(a) do Sindicato"),
    ("telefone1", "Telefone 1 (Sindicato)"),
    ("telefone2", "Telefone 2 (Sindicato)"),
    ("email1", "E-mail 1 (Sindicato)"),
    ("email2", "E-mail 2 (Sindicato)"),
    ("email3", "E-mail 3 (Sindicato)"),
    ("email4", "E-mail 4 (Sindicato)"),
    ("vice_presidente_texto", "Vice-Presidente Regional"),
    ("telefone_vice_presidente", "Telefone do Vice-Presidente Regional"),
    ("email_vice_presidente", "E-mail do Vice-Presidente Regional"),
]


def _limpar(valor):
    if valor is None:
        return None
    valor = str(valor).strip()
    return valor or None


_PALAVRAS_MINUSCULAS_TITULO = {"de", "da", "do", "das", "dos", "e"}


def _capitalizar_primeira_letra(palavra):
    """Deixa maiúscula a primeira LETRA da palavra, não o primeiro caractere
    — importante pra casos como "(marco)" virar "(Marco)" e não ficar com o
    "m" minúsculo por causa do parêntese na frente."""
    for indice, caractere in enumerate(palavra):
        if caractere.isalpha():
            return palavra[:indice] + caractere.upper() + palavra[indice + 1:]
    return palavra


def _titulo_capitalizado(texto):
    """Normaliza texto (nomes, endereços etc.) pra sempre começar cada
    palavra com maiúscula — independente de ter sido digitado TUDO EM
    MAIÚSCULO, tudo em minúsculo ou misturado. Mantém preposições comuns
    (de/da/do/das/dos/e) em minúsculo quando não são a primeira palavra,
    do jeito que normalmente se escreve em português."""
    if not texto:
        return texto
    palavras = texto.strip().lower().split(" ")
    resultado = []
    for indice, palavra in enumerate(palavras):
        if not palavra:
            resultado.append(palavra)
        elif indice > 0 and palavra in _PALAVRAS_MINUSCULAS_TITULO:
            resultado.append(palavra)
        else:
            resultado.append(_capitalizar_primeira_letra(palavra))
    return " ".join(resultado)


def listar_sindicatos(db, regiao_faec=None):
    """Lista os sindicatos (ativos primeiro, depois inativos, cada grupo em
    ordem alfabética). Se `regiao_faec` for informado, mostra só os
    sindicatos que têm pelo menos um município naquela região — assim o
    filtro de sindicato "acompanha" a região escolhida."""
    query = db.query(SindicatoRural).order_by(SindicatoRural.ativo.desc(), SindicatoRural.nome)
    if regiao_faec:
        query = (
            query.join(MunicipioSindicato, MunicipioSindicato.sindicato_id == SindicatoRural.id)
            .filter(MunicipioSindicato.regiao_faec == regiao_faec)
            .distinct()
        )
    return query.all()


def listar_regioes_faec(db):
    linhas = (
        db.query(MunicipioSindicato.regiao_faec)
        .distinct()
        .order_by(MunicipioSindicato.regiao_faec)
        .all()
    )
    return [r[0] for r in linhas if r[0]]


def listar_municipios(db, busca="", regiao_faec="", sindicato_id=None):
    """Lista municípios com filtro opcional por nome, região FAEC ou
    sindicato responsável."""
    query = db.query(MunicipioSindicato).options(joinedload(MunicipioSindicato.sindicato).defer(SindicatoRural.foto_presidente).defer(SindicatoRural.foto_presidente_tipo))

    if busca:
        query = query.filter(MunicipioSindicato.nome.ilike(f"%{busca}%"))
    if regiao_faec:
        query = query.filter(MunicipioSindicato.regiao_faec == regiao_faec)
    if sindicato_id:
        query = query.filter(MunicipioSindicato.sindicato_id == sindicato_id)

    return query.order_by(MunicipioSindicato.nome).all()


def obter_municipio(db, cod_ibge):
    municipio = db.get(MunicipioSindicato, cod_ibge)
    if municipio is not None:
        _ = municipio.sindicato  # carrega o relacionamento (join simples via lazy load)
    return municipio


def obter_historico(db, cod_ibge):
    return (
        db.query(HistoricoSindicato)
        .filter(HistoricoSindicato.cod_ibge == cod_ibge)
        .order_by(HistoricoSindicato.alterado_em.desc())
        .all()
    )


def obter_ou_criar_sindicato(db, nome):
    nome = _limpar(nome)
    if not nome:
        return None
    sindicato = db.query(SindicatoRural).filter(SindicatoRural.nome == nome).first()
    if sindicato is None:
        sindicato = SindicatoRural(nome=nome)
        db.add(sindicato)
        db.flush()
    return sindicato


def salvar_edicao(db, municipio, form, alterado_por, observacao):
    """Aplica as mudanças vindas do formulário, registrando em
    `sindicatos.historico` cada campo que realmente mudou (incluindo o
    sindicato responsável). Faz commit no final."""

    alterado_por = _limpar(alterado_por) or "não informado"
    observacao = _limpar(observacao)

    # --- Sindicato responsável (via seleção existente OU nome novo digitado) ---
    novo_nome_sindicato = _limpar(form.get("sindicato_novo"))
    if novo_nome_sindicato:
        novo_sindicato = obter_ou_criar_sindicato(db, novo_nome_sindicato)
    else:
        sindicato_id_selecionado = form.get("sindicato_id", type=int)
        novo_sindicato = (
            db.get(SindicatoRural, sindicato_id_selecionado)
            if sindicato_id_selecionado
            else None
        )

    nome_antigo = municipio.sindicato.nome if municipio.sindicato else None
    nome_novo = novo_sindicato.nome if novo_sindicato else None
    if (nome_antigo or None) != (nome_novo or None):
        db.add(HistoricoSindicato(
            cod_ibge=municipio.cod_ibge,
            campo="Sindicato Responsável",
            valor_antigo=nome_antigo,
            valor_novo=nome_novo,
            alterado_por=alterado_por,
            observacao=observacao,
        ))
        municipio.sindicato_id = novo_sindicato.id if novo_sindicato else None

    # --- Demais campos (texto simples) ---
    algo_mudou = (nome_antigo or None) != (nome_novo or None)
    for campo, rotulo in CAMPOS_EDITAVEIS.items():
        valor_novo = _limpar(form.get(campo))
        valor_antigo = getattr(municipio, campo)
        if (valor_antigo or None) != (valor_novo or None):
            db.add(HistoricoSindicato(
                cod_ibge=municipio.cod_ibge,
                campo=rotulo,
                valor_antigo=valor_antigo,
                valor_novo=valor_novo,
                alterado_por=alterado_por,
                observacao=observacao,
            ))
            setattr(municipio, campo, valor_novo)
            algo_mudou = True

    if algo_mudou:
        municipio.atualizado_em = datetime.utcnow()
        municipio.atualizado_por = alterado_por

    db.commit()
    return algo_mudou


def dados_mapa_sindicatos(db, regiao_filtro=None):
    """Monta os dados pros dois mapas do Ceará (sindicatos georreferenciados
    e municípios coloridos por região FAEC):
    - `regioes`: nomes das regiões FAEC que existem hoje, cada uma já com
      uma cor fixa (mesma região = mesma cor nos dois mapas).
    - `sindicatos`: só os que têm latitude/longitude cadastradas, com a
      região "herdada" do município mais completo vinculado a ele — se
      `regiao_filtro` for informado, só os sindicatos dessa região (o
      mapa de contorno à direita continua mostrando todas, como
      referência de onde essa região fica no estado).
    - `regiao_por_cod_ibge`: {cod_ibge: nome_da_regiao}, usado pra colorir
      o contorno de cada um dos 184 municípios no mapa de regiões."""
    paleta = [
        "#3F6F4C", "#C98A2B", "#A6412B", "#2B6CA6", "#7A4FA6",
        "#2B9E8F", "#B6482B", "#4C6E3F", "#A62B7A", "#6F6F2B",
        "#2B4CA6", "#A67A2B", "#3F8F6F", "#8F3F6F", "#5A2BA6",
    ]
    nomes_regioes = listar_regioes_faec(db)
    cor_por_regiao = {nome: paleta[i % len(paleta)] for i, nome in enumerate(nomes_regioes)}

    todos_municipios = listar_municipios_para_selecao(db)
    regiao_por_cod_ibge = {
        str(m.cod_ibge): m.regiao_faec for m in todos_municipios if m.regiao_faec
    }

    # Agrupa os municípios por sindicato UMA VEZ, em memória (reaproveitando
    # os dados que já buscamos acima) — em vez de fazer uma consulta ao
    # banco POR SINDICATO dentro do loop abaixo. Com 48+ sindicatos
    # cadastrados, isso era 48+ idas e vindas ao banco só nessa tela,
    # deixando o mapa bem lento pra carregar.
    municipios_por_sindicato_id = {}
    for m in todos_municipios:
        if m.sindicato_id:
            municipios_por_sindicato_id.setdefault(m.sindicato_id, []).append(m)

    sindicatos = []
    for sindicato in listar_sindicatos(db):
        if sindicato.latitude is None or sindicato.longitude is None:
            continue
        municipios = municipios_por_sindicato_id.get(sindicato.id, [])
        referencia = melhor_referencia(municipios)
        regiao = referencia.regiao_faec if referencia else None
        if regiao_filtro and regiao != regiao_filtro:
            continue
        sindicatos.append({
            "id": sindicato.id,
            "nome": sindicato.nome,
            "ativo": sindicato.ativo,
            "lat": float(sindicato.latitude),
            "lng": float(sindicato.longitude),
            "regiao": regiao,
            "cor": cor_por_regiao.get(regiao, "#6E6555"),
            "municipios": [m.nome for m in municipios],
        })

    return {
        "regioes": [{"nome": nome, "cor": cor_por_regiao[nome]} for nome in nomes_regioes],
        "sindicatos": sindicatos,
        "regiao_por_cod_ibge": regiao_por_cod_ibge,
        "cor_por_regiao": cor_por_regiao,
    }


CARGOS_DIRETORIA = [
    "Vice-Presidente",
    "1º Tesoureiro(a)",
    "2º Tesoureiro(a)",
    "Conselho Fiscal (titular)",
    "Conselho Fiscal (suplente)",
    "Delegado(a) representante",
]

SITUACAO_PROCESSO_LABEL = {
    "planejado": "Planejado",
    "edital_publicado": "Edital publicado",
    "registro_chapa_aberto": "Registro de chapa aberto",
    "eleicao_realizada": "Eleição realizada",
    "posse_realizada": "Posse realizada",
}
# ordem em que a situação normalmente avança, do começo pro fim do processo
SITUACAO_PROCESSO_ORDEM = list(SITUACAO_PROCESSO_LABEL.keys())


def listar_processos_eleitorais(db, sindicato_id):
    """Todo o histórico de processos eleitorais de um sindicato, do mais
    recente pro mais antigo."""
    return (
        db.query(ProcessoEleitoral)
        .filter(ProcessoEleitoral.sindicato_id == sindicato_id)
        .order_by(ProcessoEleitoral.id.desc())
        .all()
    )


def obter_processo_eleitoral(db, processo_id):
    return db.get(ProcessoEleitoral, processo_id)


def _aplicar_campos_processo(processo, form):
    processo.edital_data = _texto_para_data(form.get("edital_data", ""))
    processo.registro_chapa_prazo = _texto_para_data(form.get("registro_chapa_prazo", ""))
    processo.eleicao_data = _texto_para_data(form.get("eleicao_data", ""))
    processo.posse_data = _texto_para_data(form.get("posse_data", ""))
    processo.mandato_fim = _texto_para_data(form.get("mandato_fim", ""))
    processo.situacao = form.get("situacao") if form.get("situacao") in SITUACAO_PROCESSO_LABEL else "planejado"
    processo.presidente_eleito = _limpar(form.get("presidente_eleito", ""))
    processo.secretario_eleito = _limpar(form.get("secretario_eleito", ""))
    processo.observacoes = _limpar(form.get("observacoes", ""))


def _substituir_membros(db, processo, cargos, nomes):
    """Apaga os membros atuais do processo e recria a partir das duas listas
    paralelas vindas do formulário (um <select> de cargo + um <input> de
    nome por linha). Ignora linhas sem nome preenchido."""
    processo.membros.clear()
    ordem = 0
    for cargo, nome in zip(cargos, nomes):
        nome = _limpar(nome)
        if not nome or not cargo:
            continue
        processo.membros.append(DiretoriaMembro(cargo=cargo, nome=nome, ordem=ordem))
        ordem += 1


def _aplicar_automaticamente_se_posse_realizada(db, processo):
    """Se a situação for \"Posse realizada\" e as datas de início/fim do
    mandato estiverem preenchidas, já aplica esse processo como o mandato
    ATUAL do sindicato sozinho — sem precisar clicar em \"Usar como mandato
    atual\" à parte. Isso é o que alimenta o aviso de \"mandato vencendo\" e
    a ficha em PDF."""
    if processo.situacao == "posse_realizada" and processo.posse_data and processo.mandato_fim:
        aplicar_processo_como_mandato_atual(db, processo)


def criar_processo_eleitoral(db, sindicato_id, form, alterado_por):
    processo = ProcessoEleitoral(sindicato_id=sindicato_id)
    _aplicar_campos_processo(processo, form)
    processo.atualizado_por = _limpar(alterado_por) or "não informado"
    db.add(processo)
    db.flush()
    _substituir_membros(db, processo, form.getlist("membro_cargo"), form.getlist("membro_nome"))
    db.commit()
    _aplicar_automaticamente_se_posse_realizada(db, processo)
    return processo


def atualizar_processo_eleitoral(db, processo, form, alterado_por):
    _aplicar_campos_processo(processo, form)
    processo.atualizado_em = datetime.utcnow()
    processo.atualizado_por = _limpar(alterado_por) or "não informado"
    _substituir_membros(db, processo, form.getlist("membro_cargo"), form.getlist("membro_nome"))
    db.commit()
    _aplicar_automaticamente_se_posse_realizada(db, processo)


def excluir_processo_eleitoral(db, processo):
    db.delete(processo)
    db.commit()


def aplicar_processo_como_mandato_atual(db, processo):
    """Copia a posse e o fim do mandato deste processo pro sindicato, como o
    mandato ATUAL — é o que alimenta o aviso de \"mandato vencendo\" e a
    ficha em PDF. Também atualiza o Presidente e o Secretário(a) em TODOS os
    municípios vinculados a esse sindicato, com quem foi eleito neste
    processo (se informado) — senão a data do mandato mudava mas o nome do
    presidente continuava sendo o antigo. Acontece sozinho ao salvar um
    processo com situação \"Posse realizada\" (com as datas preenchidas);
    este botão \"Usar como mandato atual\" existe pra reaplicar manualmente
    em outros casos (ex.: reverter pra um processo mais antigo)."""
    sindicato = processo.sindicato
    sindicato.gestao_inicio = processo.posse_data
    sindicato.gestao_fim = processo.mandato_fim

    if processo.presidente_eleito or processo.secretario_eleito:
        for municipio in municipios_do_sindicato(db, sindicato.id):
            if processo.presidente_eleito:
                municipio.presidente_nome = processo.presidente_eleito
            if processo.secretario_eleito:
                municipio.secretario_nome = processo.secretario_eleito

    db.commit()


def filtrar_e_ordenar_por_mandato_vencendo(sindicatos, anos):
    """De uma lista de sindicatos, mantém só os cujo mandato atual
    (SindicatoRural.gestao_fim) vence em algum dos `anos` informados, do
    mais urgente (vence primeiro, ou já venceu) pro menos urgente. Usado
    pra destacar quem precisa de um processo eleitoral novo com prioridade,
    em vez de mostrar a lista inteira de sindicatos sem processo."""
    filtrados = [s for s in sindicatos if s.gestao_fim and s.gestao_fim.year in anos]
    filtrados.sort(key=lambda s: s.gestao_fim)
    return filtrados


def listar_sindicatos_com_processo_recente(db):
    """Todos os sindicatos com o processo eleitoral mais recente de cada um
    (ou None, se ainda não tiver nenhum cadastrado) — pra tela geral de
    Processos Eleitorais no menu, sem precisar entrar sindicato por
    sindicato pra ver a situação de cada um.

    Faz só 2 consultas ao banco no total (não 1 por sindicato) — com muitos
    sindicatos cadastrados, consultar um por um deixava a tela bem lenta."""
    sindicatos = listar_sindicatos(db)

    ultimo_id_por_sindicato = (
        db.query(
            ProcessoEleitoral.sindicato_id,
            func.max(ProcessoEleitoral.id).label("ultimo_id"),
        )
        .group_by(ProcessoEleitoral.sindicato_id)
        .subquery()
    )
    processos_recentes = (
        db.query(ProcessoEleitoral)
        .join(ultimo_id_por_sindicato, ProcessoEleitoral.id == ultimo_id_por_sindicato.c.ultimo_id)
        .options(joinedload(ProcessoEleitoral.membros))
        .all()
    )
    processo_por_sindicato_id = {p.sindicato_id: p for p in processos_recentes}

    return [(s, processo_por_sindicato_id.get(s.id)) for s in sindicatos]


def gerar_planilha_processos_eleitorais_xlsx(db):
    """Gera uma planilha .xlsx com a situação do processo eleitoral mais
    recente de cada sindicato — pronta pra distribuir/imprimir. Retorna
    bytes prontos para download."""
    sindicatos_processos = listar_sindicatos_com_processo_recente(db)

    wb = Workbook()
    ws = wb.active
    ws.title = "Processos_Eleitorais"

    colunas = [
        "Sindicato", "Situação", "Edital de Convocação", "Prazo de Registro de Chapa",
        "Data da Eleição", "Início do Mandato (Posse)", "Fim do Mandato",
        "Presidente Eleito", "Secretário(a) Eleito(a)", "Demais membros da diretoria",
        "Observações",
    ]
    ws.append(colunas)
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="3F6F4C")

    for sindicato, processo in sindicatos_processos:
        if processo is None:
            ws.append([sindicato.nome, "Nenhum processo cadastrado ainda", None, None, None, None, None, None, None, None, None])
            continue
        membros_txt = "; ".join(f"{m.cargo}: {m.nome}" for m in processo.membros)
        ws.append([
            sindicato.nome,
            SITUACAO_PROCESSO_LABEL.get(processo.situacao, processo.situacao),
            processo.edital_data,
            processo.registro_chapa_prazo,
            processo.eleicao_data,
            processo.posse_data,
            processo.mandato_fim,
            processo.presidente_eleito,
            processo.secretario_eleito,
            membros_txt or None,
            processo.observacoes,
        ])

    for indice, rotulo in enumerate(colunas, start=1):
        ws.column_dimensions[get_column_letter(indice)].width = max(16, len(rotulo) + 2)
    for linha in ws.iter_rows(min_row=2, min_col=3, max_col=7):
        for celula in linha:
            if celula.value is not None:
                celula.number_format = "DD/MM/YYYY"

    ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()


# Pro "Cronograma Consolidado de Eleições Sindicais": agrupa a situação
# detalhada em só 3 status/cores, do jeito que já era feito no controle em
# planilha manual (verde = já votou, amarelo = processo em andamento,
# cinza = ainda nem começou).
_STATUS_CRONOGRAMA_POR_SITUACAO = {
    None: ("A Iniciar (Aguardando Edital)", "cinza"),
    "planejado": ("A Iniciar (Aguardando Edital)", "cinza"),
    "edital_publicado": ("Edital Publicado (Em Andamento)", "amarelo"),
    "registro_chapa_aberto": ("Edital Publicado (Em Andamento)", "amarelo"),
    "eleicao_realizada": ("Eleição Realizada", "verde"),
    "posse_realizada": ("Eleição Realizada", "verde"),
}
_CORES_CRONOGRAMA_XLSX = {
    "verde": "D6E9D0",
    "amarelo": "FCEFC7",
    "cinza": "E7E3D8",
}


def gerar_cronograma_eleitoral_xlsx(db):
    """Gera a planilha \"Cronograma Consolidado de Eleições Sindicais\": uma
    linha por sindicato com início/término do mandato atual e as datas do
    processo eleitoral mais recente (edital, registro de chapa, eleição,
    posse), com as linhas coloridas por status (verde = eleição já
    realizada, amarelo = edital publicado/processo em andamento, cinza =
    ainda aguardando o edital) — ordenado por término do mandato, do que
    vence primeiro pro que vence por último. Retorna bytes prontos pra
    download."""
    sindicatos_processos = listar_sindicatos_com_processo_recente(db)

    def chave_ordenacao(par):
        sindicato, _ = par
        return sindicato.gestao_fim or date.max

    sindicatos_processos = sorted(sindicatos_processos, key=chave_ordenacao)

    wb = Workbook()
    ws = wb.active
    ws.title = "Cronograma_Eleicoes"

    ws.append(["CRONOGRAMA CONSOLIDADO DE ELEIÇÕES SINDICAIS"])
    ws.merge_cells("A1:J1")
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Verde: Eleição Realizada  |  Amarelo: Edital Publicado (Em Andamento)  |  Cinza: A Iniciar (Aguardando Edital)"])
    ws.merge_cells("A2:J2")
    ws["A2"].font = Font(italic=True, size=9, color="6E6555")
    ws.append([])

    cabecalho = [
        "Nº", "Sindicato", "Início Mandato", "Término Mandato",
        "Edital Eleição", "Edital Reg. Chapa", "Data Eleição", "Data Posse",
        "Status do Processo",
    ]
    ws.append(cabecalho)
    linha_cabecalho = ws.max_row
    for celula in ws[linha_cabecalho]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="3F6F4C")

    for indice, (sindicato, processo) in enumerate(sindicatos_processos, start=1):
        rotulo_status, grupo_cor = _STATUS_CRONOGRAMA_POR_SITUACAO.get(
            processo.situacao if processo else None, ("A Iniciar (Aguardando Edital)", "cinza")
        )
        ws.append([
            indice,
            sindicato.nome,
            sindicato.gestao_inicio,
            sindicato.gestao_fim,
            processo.edital_data if processo else None,
            processo.registro_chapa_prazo if processo else None,
            processo.eleicao_data if processo else None,
            processo.posse_data if processo else None,
            rotulo_status,
        ])
        linha = ws.max_row
        cor_fundo = _CORES_CRONOGRAMA_XLSX[grupo_cor]
        for celula in ws[linha]:
            celula.fill = PatternFill("solid", fgColor=cor_fundo)
        for coluna in (3, 4, 5, 6, 7, 8):
            celula = ws.cell(row=linha, column=coluna)
            if celula.value is not None:
                celula.number_format = "DD/MM/YYYY"

    larguras = [5, 24, 14, 14, 14, 16, 13, 13, 30]
    for indice, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(indice)].width = largura
    ws.freeze_panes = f"A{linha_cabecalho + 1}"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()


def obter_sindicato(db, sindicato_id):
    return db.get(SindicatoRural, sindicato_id)


def melhor_referencia(municipios):
    """Escolhe, entre os municípios de um sindicato, o que tem MAIS campos
    de contato preenchidos, pra usar como referência ao mostrar/editar os
    dados do sindicato. Assim, remover ou adicionar um município não faz o
    formulário "esquecer" informações que só apareciam em outro município do
    mesmo sindicato — antes disso, sempre pegava o primeiro da lista, que
    podia estar com campos em branco."""
    if not municipios:
        return None
    def completude(m):
        return sum(1 for campo in CAMPOS_EDITAVEIS if getattr(m, campo))
    return max(municipios, key=completude)


def desvincular_municipio(db, municipio, alterado_por, observacao=None):
    """Remove o vínculo deste município com o sindicato atual dele — fica
    'sem sindicato responsável' até alguém vincular a outro. Antes de tirar o
    vínculo, se este município tinha dados de contato (presidente,
    secretário, telefones, e-mails, vice-presidente) que os outros do MESMO
    sindicato não tinham, esses dados são copiados pros outros primeiro —
    assim a informação continua visível na tela do sindicato mesmo depois
    de remover justamente o município que estava "guardando" ela.
    Registra no histórico. Retorna False se ele já não tinha sindicato."""
    alterado_por = _limpar(alterado_por) or "não informado"
    observacao = _limpar(observacao) or "Removido do sindicato pelo painel"

    if not municipio.sindicato:
        return False
    nome_antigo = municipio.sindicato.nome
    sindicato_id_antigo = municipio.sindicato_id

    outros = [m for m in municipios_do_sindicato(db, sindicato_id_antigo) if m.cod_ibge != municipio.cod_ibge]
    for outro in outros:
        for campo in CAMPOS_EDITAVEIS:
            valor_daqui = getattr(municipio, campo)
            if valor_daqui and not getattr(outro, campo):
                setattr(outro, campo, valor_daqui)
                outro.atualizado_em = datetime.utcnow()
                outro.atualizado_por = alterado_por

    db.add(HistoricoSindicato(
        cod_ibge=municipio.cod_ibge, campo="Sindicato Responsável",
        valor_antigo=nome_antigo, valor_novo=None,
        alterado_por=alterado_por, observacao=observacao,
    ))
    municipio.sindicato_id = None
    municipio.atualizado_em = datetime.utcnow()
    municipio.atualizado_por = alterado_por
    db.commit()
    return True


def vincular_municipio_a_sindicato(db, municipio, sindicato, alterado_por, observacao=None, copiar_contato=True):
    """Vincula este município a um sindicato já existente (tira de onde
    estava, se estava em outro). Se `copiar_contato`, copia os dados de
    contato (presidente, secretário, vice-presidente, telefones, e-mails) de
    outro município que já está nesse sindicato, pra manter tudo igual entre
    eles — do mesmo jeito que já acontece quando o sindicato é criado com
    vários municípios de uma vez. Retorna False se ele já estava lá."""
    alterado_por = _limpar(alterado_por) or "não informado"
    observacao = _limpar(observacao) or f'Vinculado ao sindicato "{sindicato.nome}" pelo painel'

    nome_antigo = municipio.sindicato.nome if municipio.sindicato else None
    if nome_antigo == sindicato.nome:
        return False

    db.add(HistoricoSindicato(
        cod_ibge=municipio.cod_ibge, campo="Sindicato Responsável",
        valor_antigo=nome_antigo, valor_novo=sindicato.nome,
        alterado_por=alterado_por, observacao=observacao,
    ))
    municipio.sindicato_id = sindicato.id

    if copiar_contato:
        outros = [m for m in municipios_do_sindicato(db, sindicato.id) if m.cod_ibge != municipio.cod_ibge]
        referencia = melhor_referencia(outros)
        if referencia:
            for campo, rotulo in CAMPOS_EDITAVEIS.items():
                valor_novo = getattr(referencia, campo)
                valor_antigo = getattr(municipio, campo)
                if (valor_antigo or None) != (valor_novo or None):
                    db.add(HistoricoSindicato(
                        cod_ibge=municipio.cod_ibge, campo=rotulo,
                        valor_antigo=valor_antigo, valor_novo=valor_novo,
                        alterado_por=alterado_por, observacao=observacao,
                    ))
                    setattr(municipio, campo, valor_novo)

    municipio.atualizado_em = datetime.utcnow()
    municipio.atualizado_por = alterado_por
    db.commit()
    return True


def atualizar_status_sindicato(db, sindicato, ativo):
    """Marca um sindicato como ativo/inativo. Usado quando ele é substituído
    por outro (municípios reatribuídos) mas ainda deve continuar existindo
    no histórico, só que sinalizado como não estando mais em atividade."""
    if sindicato.ativo != ativo:
        sindicato.ativo = ativo
        db.commit()
        return True
    return False


def renomear_sindicato(db, sindicato, novo_nome):
    """Troca o nome de um sindicato existente (ex.: correção de digitação,
    ou o sindicato passou a se chamar outra coisa). Levanta ValueError se já
    existir outro sindicato com esse nome."""
    novo_nome = _limpar(novo_nome)
    if not novo_nome:
        raise ValueError("O nome do sindicato não pode ficar em branco.")
    if novo_nome.lower() == (sindicato.nome or "").lower():
        return False
    existente = (
        db.query(SindicatoRural)
        .filter(SindicatoRural.nome.ilike(novo_nome), SindicatoRural.id != sindicato.id)
        .first()
    )
    if existente:
        raise ValueError(f'Já existe outro sindicato chamado "{existente.nome}".')
    sindicato.nome = novo_nome
    db.commit()
    return True


FOTO_TIPOS_PERMITIDOS = {
    "image/jpeg": "jpg", "image/jpg": "jpg", "image/png": "png", "image/webp": "webp",
}
FOTO_TAMANHO_MAXIMO_BYTES = 5 * 1024 * 1024  # 5 MB


def _texto_para_float(valor):
    valor = _limpar(valor)
    if not valor:
        return None
    try:
        return float(valor.replace(",", "."))
    except ValueError:
        return None


def _texto_para_data(valor):
    valor = _limpar(valor)
    if not valor:
        return None
    try:
        return datetime.strptime(valor, "%Y-%m-%d").date()
    except ValueError:
        return None


def atualizar_endereco_geo_foto_sindicato(db, sindicato, endereco, latitude, longitude, arquivo_foto,
                                           gestao_inicio=None, gestao_fim=None):
    """Atualiza endereço, coordenadas (georreferência), início/fim da gestão
    do presidente atual e, se enviada, a foto do presidente de um sindicato.
    A foto vai pro Supabase Storage (não mais bytea na tabela — isso é o que
    estava inflando o egress em toda listagem que faz joinedload no
    sindicato). `arquivo_foto` é o FileStorage do Flask
    (request.files.get(...)) ou None. Levanta ValueError se o arquivo não for
    uma imagem válida ou se o Storage não estiver configurado."""

    sindicato.endereco = _limpar(endereco)
    sindicato.latitude = _texto_para_float(latitude)
    sindicato.longitude = _texto_para_float(longitude)
    sindicato.gestao_inicio = _texto_para_data(gestao_inicio)
    sindicato.gestao_fim = _texto_para_data(gestao_fim)

    if arquivo_foto and arquivo_foto.filename:
        tipo = (arquivo_foto.mimetype or "").lower()
        if tipo not in FOTO_TIPOS_PERMITIDOS:
            raise ValueError("A foto precisa ser um arquivo JPG, PNG ou WEBP.")
        conteudo = arquivo_foto.read()
        if len(conteudo) > FOTO_TAMANHO_MAXIMO_BYTES:
            raise ValueError("A foto passou de 5 MB — reduza o tamanho e tente de novo.")

        url_antiga = sindicato.foto_presidente_url
        try:
            nova_url = supabase_storage.upload_foto_presidente(
                sindicato.id, conteudo, tipo, arquivo_foto.filename
            )
        except supabase_storage.SupabaseStorageError as erro:
            raise ValueError(f"Não consegui enviar a foto pro Storage: {erro}")

        sindicato.foto_presidente_url = nova_url
        # Limpa os campos antigos (se essa linha ainda tinha foto guardada
        # em binário de antes da migração) pra não ficar bytea morto na tabela.
        sindicato.foto_presidente = None
        sindicato.foto_presidente_tipo = None

        if url_antiga:
            supabase_storage.excluir_arquivo_por_url(url_antiga)

    db.commit()


def sindicatos_com_gestao_vencendo(db, ate=None):
    """Sindicatos ativos cuja gestão do presidente tem data de fim marcada e
    está a caminho de vencer (ou já venceu). Por padrão, `ate` é o que vier
    depois entre 31/12 do ano atual e hoje+90 dias — ou seja, mostra todo
    mundo que vence este ano E também garante pelo menos 90 dias de
    antecedência (importante nos últimos meses do ano, pra não "sumir" um
    mandato que vence em janeiro/fevereiro do ano seguinte só porque o ano
    ainda não virou). Usado pra avisar dentro do próprio sistema, sem
    precisar abrir a ficha em PDF pra descobrir. Retorna uma lista de
    (sindicato, dias_restantes), ordenada do mais urgente pro menos urgente
    (vencidos primeiro)."""
    hoje = date.today()
    limite = ate or max(date(hoje.year, 12, 31), hoje + timedelta(days=90))

    sindicatos = (
        db.query(SindicatoRural)
        .filter(SindicatoRural.ativo.is_(True))
        .filter(SindicatoRural.gestao_fim.isnot(None))
        .filter(SindicatoRural.gestao_fim <= limite)
        .all()
    )

    resultado = [(s, (s.gestao_fim - hoje).days) for s in sindicatos]
    resultado.sort(key=lambda par: par[1])
    return resultado


def criar_sindicato(db, nome):
    """Cadastra um sindicato novo, sem vincular nenhum município ainda.
    Levanta ValueError se já existir um sindicato com esse nome."""
    nome = _limpar(nome)
    if not nome:
        raise ValueError("Informe um nome para o sindicato.")
    existente = db.query(SindicatoRural).filter(SindicatoRural.nome.ilike(nome)).first()
    if existente:
        raise ValueError(f'Já existe um sindicato chamado "{existente.nome}".')
    sindicato = SindicatoRural(nome=nome, ativo=True)
    db.add(sindicato)
    db.commit()
    return sindicato


def listar_municipios_para_selecao(db):
    """Todos os municípios (nome + sindicato atual), pra popular o seletor
    de "quais municípios entram nesse sindicato novo"."""
    return (
        db.query(MunicipioSindicato)
        .options(joinedload(MunicipioSindicato.sindicato).defer(SindicatoRural.foto_presidente).defer(SindicatoRural.foto_presidente_tipo))
        .order_by(MunicipioSindicato.nome)
        .all()
    )


def criar_sindicato_com_municipios(db, nome, form, cod_ibges_selecionados, alterado_por, observacao):
    """Cadastra o sindicato E já atribui os municípios escolhidos a ele,
    aplicando os mesmos dados de contato do formulário a cada um. Levanta
    ValueError se já existir um sindicato com esse nome."""
    sindicato = criar_sindicato(db, nome)

    alterado_por = _limpar(alterado_por) or "não informado"
    observacao = _limpar(observacao) or f'Cadastro do sindicato "{sindicato.nome}"'
    novos_valores = {campo: _limpar(form.get(campo)) for campo in CAMPOS_EDITAVEIS}

    total = 0
    for cod_ibge in cod_ibges_selecionados:
        m = db.get(MunicipioSindicato, cod_ibge)
        if not m:
            continue

        nome_sindicato_antigo = m.sindicato.nome if m.sindicato else None
        if nome_sindicato_antigo != sindicato.nome:
            db.add(HistoricoSindicato(
                cod_ibge=m.cod_ibge, campo="Sindicato Responsável",
                valor_antigo=nome_sindicato_antigo, valor_novo=sindicato.nome,
                alterado_por=alterado_por, observacao=observacao,
            ))
            m.sindicato_id = sindicato.id

        for campo, rotulo in CAMPOS_EDITAVEIS.items():
            valor_novo = novos_valores[campo]
            valor_antigo = getattr(m, campo)
            if (valor_antigo or None) != (valor_novo or None):
                db.add(HistoricoSindicato(
                    cod_ibge=m.cod_ibge, campo=rotulo,
                    valor_antigo=valor_antigo, valor_novo=valor_novo,
                    alterado_por=alterado_por, observacao=observacao,
                ))
                setattr(m, campo, valor_novo)

        m.atualizado_em = datetime.utcnow()
        m.atualizado_por = alterado_por
        total += 1

    db.commit()
    return sindicato, total


def municipios_do_sindicato(db, sindicato_id):
    """Todos os municípios vinculados hoje a este sindicato (usado tanto pra
    mostrar a lista quanto pra aplicar uma edição em lote)."""
    return (
        db.query(MunicipioSindicato)
        .filter(MunicipioSindicato.sindicato_id == sindicato_id)
        .order_by(MunicipioSindicato.nome)
        .all()
    )


def salvar_edicao_em_lote_sindicato(db, sindicato, form, alterado_por, observacao):
    """Aplica os mesmos dados de contato (presidente, vice-presidente,
    telefones, e-mails) a TODOS os municípios vinculados a este sindicato de
    uma vez só — pra não precisar editar município por município quando só o
    presidente do sindicato mudou. Registra o histórico normalmente, um
    lançamento por município que realmente mudou."""

    alterado_por = _limpar(alterado_por) or "não informado"
    observacao = _limpar(observacao) or f'Alteração em lote via o sindicato "{sindicato.nome}"'

    municipios = municipios_do_sindicato(db, sindicato.id)
    novos_valores = {campo: _limpar(form.get(campo)) for campo in CAMPOS_EDITAVEIS}

    total_alterados = 0
    for m in municipios:
        mudou_este = False
        for campo, rotulo in CAMPOS_EDITAVEIS.items():
            valor_novo = novos_valores[campo]
            valor_antigo = getattr(m, campo)
            if (valor_antigo or None) != (valor_novo or None):
                db.add(HistoricoSindicato(
                    cod_ibge=m.cod_ibge, campo=rotulo,
                    valor_antigo=valor_antigo, valor_novo=valor_novo,
                    alterado_por=alterado_por, observacao=observacao,
                ))
                setattr(m, campo, valor_novo)
                mudou_este = True
        if mudou_este:
            m.atualizado_em = datetime.utcnow()
            m.atualizado_por = alterado_por
            total_alterados += 1

    db.commit()
    return total_alterados, len(municipios)


def gerar_planilha_xlsx(db, busca="", regiao_faec="", sindicato_id=None):
    """Gera, a partir do estado ATUAL do banco, uma planilha .xlsx com as
    mesmas colunas da planilha original — pronta para distribuir aos
    setores. Respeita os mesmos filtros aplicados na tela (busca, região,
    sindicato); sem filtro nenhum, exporta todos os municípios. Retorna
    bytes prontos para download."""

    municipios = listar_municipios(db, busca, regiao_faec, sindicato_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Municipios_Sindicatos"

    cabecalho = [rotulo for _, rotulo in COLUNAS_EXPORTACAO]
    ws.append(cabecalho)
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="3F6F4C")

    for m in municipios:
        linha = []
        for atributo, _ in COLUNAS_EXPORTACAO:
            if atributo == "_sindicato_nome":
                linha.append(m.sindicato.nome if m.sindicato else None)
            else:
                linha.append(getattr(m, atributo))
        ws.append(linha)

    for indice, (atributo, rotulo) in enumerate(COLUNAS_EXPORTACAO, start=1):
        largura = max(14, len(rotulo) + 2)
        ws.column_dimensions[get_column_letter(indice)].width = largura

    ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()


# =========================================================
# Ficha "Proposta de Distribuição dos Municípios nos Sindicatos"
# Mesmo layout usado no papel: por Região FAEC, com o(a)
# Vice-Presidente Regional, e a tabela Sindicato/Presidente x Municípios.
# Gerada sempre a partir do estado ATUAL do banco.
# =========================================================

_FICHA_VERDE = colors.HexColor("#8CC63F")
_FICHA_CINZA = colors.HexColor("#B9B9B9")
_FICHA_LINHA = colors.HexColor("#9A9A9A")

_ficha_titulo_style = ParagraphStyle(
    "FichaTitulo", fontName="Helvetica-Bold", fontSize=15,
    alignment=TA_CENTER, leading=18,
)
_ficha_data_style = ParagraphStyle(
    "FichaData", fontName="Helvetica", fontSize=8.5, alignment=TA_RIGHT,
)
_ficha_regiao_style = ParagraphStyle(
    "FichaRegiao", fontName="Helvetica-Bold", fontSize=11,
    alignment=TA_CENTER, leading=14,
)
_ficha_vice_style = ParagraphStyle(
    "FichaVice", fontName="Helvetica", fontSize=8.5, alignment=TA_CENTER,
)
_ficha_cabecalho_style = ParagraphStyle(
    "FichaCabecalhoCol", fontName="Helvetica-Bold", fontSize=8.5,
    alignment=TA_CENTER,
)
_ficha_sindicato_style = ParagraphStyle(
    "FichaSindicato", fontName="Helvetica-Bold", fontSize=8.5, leading=10.5,
)
_ficha_municipios_style = ParagraphStyle(
    "FichaMunicipios", fontName="Helvetica", fontSize=8.5, leading=10.5,
)

_FICHA_COL1 = 55 * mm
_FICHA_COL2 = 190 * mm

# Larguras do "Relatório de Contatos" (3 colunas: Sindicato/Presidente/
# Telefone/E-mail) — soma igual à largura total da Ficha de Distribuição
# acima, só dividida diferente.
_RELATORIO_LARGURA_TOTAL = _FICHA_COL1 + _FICHA_COL2
_RELATORIO_COLS = [95 * mm, 35 * mm, 115 * mm]


def _limpar_texto_vice(texto):
    """O campo vice_presidente_texto já costuma começar com algo como
    "Vice-Presidente Regional da FAEC – NOME..." ou "Diretor Regional da
    FAEC: NOME...". Tira esse prefixo pra não repetir o rótulo que a ficha
    já mostra antes ("VICE-PRESIDENTE REGIONAL – ...")."""
    if not texto:
        return texto
    texto = re.sub(
        r'^\s*(vice-presidente regional( da faec)?|diretor(a)? regional( da faec)?)\s*[:\-–]\s*',
        '', texto, flags=re.IGNORECASE,
    )
    return texto.strip()


def montar_estrutura_ficha(db):
    """Agrupa os municípios atuais por Região FAEC -> Sindicato, na ordem:
    região (alfabética), sindicato (alfabética), município (alfabética)."""
    municipios = (
        db.query(MunicipioSindicato)
        .options(joinedload(MunicipioSindicato.sindicato).defer(SindicatoRural.foto_presidente).defer(SindicatoRural.foto_presidente_tipo))
        .order_by(MunicipioSindicato.regiao_faec, MunicipioSindicato.nome)
        .all()
    )

    regioes = {}
    for m in municipios:
        nome_regiao = m.regiao_faec or "Região não definida"
        regiao = regioes.setdefault(nome_regiao, {"vice_presidente": None, "sindicatos": {}})
        if not regiao["vice_presidente"] and m.vice_presidente_texto:
            regiao["vice_presidente"] = _limpar_texto_vice(m.vice_presidente_texto)

        nome_sindicato = m.sindicato.nome if m.sindicato else "Sem sindicato definido"
        info_sindicato = regiao["sindicatos"].setdefault(
            nome_sindicato, {"presidente": m.presidente_nome, "municipios": []}
        )
        if not info_sindicato["presidente"] and m.presidente_nome:
            info_sindicato["presidente"] = m.presidente_nome
        info_sindicato["municipios"].append(m.nome)

    return regioes


def _ficha_faixa_regiao(nome_regiao, vice_presidente):
    linhas = [[Paragraph(nome_regiao.upper(), _ficha_regiao_style)]]
    if vice_presidente:
        linhas.append([Paragraph(f"VICE-PRESIDENTE REGIONAL – {vice_presidente}", _ficha_vice_style)])
    t = Table(linhas, colWidths=[_FICHA_COL1 + _FICHA_COL2])
    estilo = [
        ("BACKGROUND", (0, 0), (-1, -1), _FICHA_CINZA),
        ("TOPPADDING", (0, 0), (-1, 0), 5),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 1 if vice_presidente else 5),
    ]
    if vice_presidente:
        estilo += [("TOPPADDING", (0, 1), (-1, 1), 0), ("BOTTOMPADDING", (0, 1), (-1, 1), 5)]
    t.setStyle(TableStyle(estilo))
    return t


def _ficha_tabela_regiao(sindicatos_dict):
    dados = [[
        Paragraph("SINDICATOS/PRESIDENTES", _ficha_cabecalho_style),
        Paragraph("MUNICÍPIOS", _ficha_cabecalho_style),
    ]]
    for nome_sindicato in sorted(sindicatos_dict.keys()):
        info = sindicatos_dict[nome_sindicato]
        rotulo = nome_sindicato
        if info["presidente"]:
            rotulo += f" ({info['presidente']})"
        municipios_txt = " - ".join(info["municipios"])
        dados.append([
            Paragraph(rotulo, _ficha_sindicato_style),
            Paragraph(municipios_txt, _ficha_municipios_style),
        ])

    t = Table(dados, colWidths=[_FICHA_COL1, _FICHA_COL2], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F2F2")),
        ("GRID", (0, 0), (-1, -1), 0.6, _FICHA_LINHA),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]))
    return t


def gerar_ficha_distribuicao_pdf(db):
    """Gera a "Proposta/Ficha de Distribuição dos Municípios nos Sindicatos"
    em PDF, a partir do estado ATUAL do banco (mesmo layout usado no papel:
    faixa verde de título, cada Região FAEC com uma faixa cinza + tabela
    Sindicato/Presidente x Municípios). Retorna bytes prontos para download."""

    regioes = montar_estrutura_ficha(db)

    story = []
    cabecalho = Table(
        [[Paragraph("PROPOSTA DE DISTRIBUIÇÃO DOS MUNICÍPIOS<br/>NOS SINDICATOS RURAIS", _ficha_titulo_style)]],
        colWidths=[_FICHA_COL1 + _FICHA_COL2],
    )
    cabecalho.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), _FICHA_VERDE),
        ("TOPPADDING", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
    ]))
    story.append(cabecalho)
    story.append(Spacer(1, 2 * mm))
    story.append(Paragraph(f"DATA DA ATUALIZAÇÃO: {datetime.now():%d.%m.%Y}", _ficha_data_style))
    story.append(Spacer(1, 4 * mm))

    for nome_regiao in sorted(regioes.keys()):
        info_regiao = regioes[nome_regiao]
        story.append(KeepTogether([
            _ficha_faixa_regiao(nome_regiao, info_regiao["vice_presidente"]),
            _ficha_tabela_regiao(info_regiao["sindicatos"]),
            Spacer(1, 4 * mm),
        ]))

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        topMargin=12 * mm, bottomMargin=12 * mm,
        leftMargin=12 * mm, rightMargin=12 * mm,
        title="Proposta de Distribuição dos Municípios nos Sindicatos Rurais",
    )
    doc.build(story)
    buffer.seek(0)
    return buffer.read()


def montar_estrutura_relatorio_contatos(db):
    """Agrupa os municípios por Região FAEC -> Sindicato, com os dados de
    contato de cada sindicato (presidente, telefone, e-mail) — pro
    'Relatório de Contatos dos Sindicatos por Região': uma linha por
    SINDICATO (não repete por município)."""
    municipios = (
        db.query(MunicipioSindicato)
        .options(joinedload(MunicipioSindicato.sindicato).defer(SindicatoRural.foto_presidente).defer(SindicatoRural.foto_presidente_tipo))
        .order_by(MunicipioSindicato.regiao_faec, MunicipioSindicato.nome)
        .all()
    )

    regioes = {}
    for m in municipios:
        nome_regiao = m.regiao_faec or "Região não definida"
        regiao = regioes.setdefault(nome_regiao, {"vice_presidente": None, "telefone_vice": None, "sindicatos": {}})
        if not regiao["vice_presidente"] and m.vice_presidente_texto:
            regiao["vice_presidente"] = _limpar_texto_vice(m.vice_presidente_texto)
        if not regiao["telefone_vice"] and m.telefone_vice_presidente:
            regiao["telefone_vice"] = m.telefone_vice_presidente

        nome_sindicato = m.sindicato.nome if m.sindicato else "Sem sindicato definido"
        info_sindicato = regiao["sindicatos"].setdefault(
            nome_sindicato, {"presidente": None, "telefone": None, "email": None}
        )
        if not info_sindicato["presidente"] and m.presidente_nome:
            info_sindicato["presidente"] = m.presidente_nome
        if not info_sindicato["telefone"] and m.telefone1:
            info_sindicato["telefone"] = m.telefone1
        if not info_sindicato["email"]:
            emails = [e for e in [m.email1, m.email2, m.email3, m.email4] if e]
            if emails:
                info_sindicato["email"] = " / ".join(emails)

    return regioes


def _relatorio_faixa_regiao(nome_regiao, vice_presidente, telefone_vice):
    linhas = [[Paragraph(nome_regiao.upper(), _ficha_regiao_style)]]
    if vice_presidente:
        texto_vice = f"VICE-PRESIDENTE REGIONAL – {vice_presidente}"
        if telefone_vice:
            texto_vice += f" – telefone {telefone_vice}"
        linhas.append([Paragraph(texto_vice, _ficha_vice_style)])
    t = Table(linhas, colWidths=[_RELATORIO_LARGURA_TOTAL])
    estilo = [
        ("BACKGROUND", (0, 0), (-1, -1), _FICHA_CINZA),
        ("TOPPADDING", (0, 0), (-1, 0), 5),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 1 if vice_presidente else 5),
    ]
    if vice_presidente:
        estilo += [("TOPPADDING", (0, 1), (-1, 1), 0), ("BOTTOMPADDING", (0, 1), (-1, 1), 5)]
    t.setStyle(TableStyle(estilo))
    return t


def _relatorio_tabela_contatos(sindicatos_dict):
    dados = [[
        Paragraph("SINDICATO/PRESIDENTE", _ficha_cabecalho_style),
        Paragraph("TELEFONE", _ficha_cabecalho_style),
        Paragraph("E-MAIL", _ficha_cabecalho_style),
    ]]
    for nome_sindicato in sorted(sindicatos_dict.keys()):
        info = sindicatos_dict[nome_sindicato]
        rotulo = nome_sindicato
        if info["presidente"]:
            rotulo += f" ({info['presidente']})"
        dados.append([
            Paragraph(rotulo, _ficha_sindicato_style),
            Paragraph(info["telefone"] or "—", _ficha_municipios_style),
            Paragraph(info["email"] or "—", _ficha_municipios_style),
        ])

    t = Table(dados, colWidths=_RELATORIO_COLS, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F2F2")),
        ("GRID", (0, 0), (-1, -1), 0.6, _FICHA_LINHA),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]))
    return t


def gerar_relatorio_contatos_pdf(db):
    """Gera o "Relatório de Contatos dos Sindicatos por Região FAEC" em
    PDF, a partir do estado ATUAL do banco — uma linha por SINDICATO
    (Presidente/Telefone/E-mail), agrupado por Região FAEC com o
    Vice-Presidente Regional no topo. Retorna bytes prontos para download."""
    regioes = montar_estrutura_relatorio_contatos(db)

    story = []
    cabecalho = Table(
        [[Paragraph("CONTATOS DOS SINDICATOS RURAIS<br/>POR REGIÃO FAEC", _ficha_titulo_style)]],
        colWidths=[_RELATORIO_LARGURA_TOTAL],
    )
    cabecalho.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), _FICHA_VERDE),
        ("TOPPADDING", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
    ]))
    story.append(cabecalho)
    story.append(Spacer(1, 2 * mm))
    story.append(Paragraph(f"DATA DA ATUALIZAÇÃO: {datetime.now():%d.%m.%Y}", _ficha_data_style))
    story.append(Spacer(1, 4 * mm))

    for nome_regiao in sorted(regioes.keys()):
        info_regiao = regioes[nome_regiao]
        story.append(_relatorio_faixa_regiao(nome_regiao, info_regiao["vice_presidente"], info_regiao["telefone_vice"]))
        story.append(_relatorio_tabela_contatos(info_regiao["sindicatos"]))
        story.append(Spacer(1, 4 * mm))

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        topMargin=12 * mm, bottomMargin=12 * mm,
        leftMargin=12 * mm, rightMargin=12 * mm,
        title="Contatos dos Sindicatos Rurais por Região FAEC",
    )
    doc.build(story)
    buffer.seek(0)
    return buffer.read()


# =========================================================
# Ficha individual do sindicato (dados + histórico) em PDF
# =========================================================

def obter_historico_sindicato(db, sindicato):
    """Histórico relevante a ESTE sindicato: mudanças de "Sindicato
    Responsável" que citam o nome dele (entrando ou saindo de algum
    município) + mudanças de presidente/contatos nos municípios vinculados a
    ele hoje. Como uma edição em lote grava uma linha por município, entradas
    idênticas (mesmo campo/de/para/quem/quando) são mostradas uma única vez."""

    cod_ibges_atuais = [m.cod_ibge for m in municipios_do_sindicato(db, sindicato.id)]

    eventos_nome = (
        db.query(HistoricoSindicato)
        .filter(
            HistoricoSindicato.campo == "Sindicato Responsável",
            or_(
                HistoricoSindicato.valor_novo == sindicato.nome,
                HistoricoSindicato.valor_antigo == sindicato.nome,
            ),
        )
        .all()
    )

    eventos_contatos = []
    if cod_ibges_atuais:
        eventos_contatos = (
            db.query(HistoricoSindicato)
            .filter(
                HistoricoSindicato.cod_ibge.in_(cod_ibges_atuais),
                HistoricoSindicato.campo != "Sindicato Responsável",
            )
            .all()
        )

    vistos = set()
    resultado = []
    for h in eventos_nome + eventos_contatos:
        chave = (h.campo, h.valor_antigo, h.valor_novo, h.alterado_por, h.alterado_em)
        if chave not in vistos:
            vistos.add(chave)
            resultado.append(h)

    resultado.sort(key=lambda h: h.alterado_em or datetime.min, reverse=True)
    return resultado


def _imagem_pdf_da_foto(foto_bytes, largura_max_mm=32, altura_max_mm=40):
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image as RLImage

    largura_max = largura_max_mm * mm
    altura_max = altura_max_mm * mm
    leitor = ImageReader(BytesIO(foto_bytes))
    largura_original, altura_original = leitor.getSize()
    escala = min(largura_max / float(largura_original), altura_max / float(altura_original))
    return RLImage(BytesIO(foto_bytes), width=largura_original * escala, height=altura_original * escala)


def _texto_situacao_mandato(sindicato):
    """Retorna (texto, cor_hex) descrevendo a situação do mandato do
    presidente, pra destacar na ficha quando estiver perto de vencer ou já
    vencido. cor_hex é None quando não há nada a avisar."""
    if not sindicato.gestao_fim:
        return None, None
    dias = (sindicato.gestao_fim - date.today()).days
    if dias < 0:
        return f"MANDATO VENCIDO há {-dias} dia(s)", "#C0392B"
    if dias <= 60:
        return f"vence em {dias} dia(s)", "#B8790A"
    return None, None


def _chips_municipios(municipios, por_linha=3, largura_total_mm=180):
    """Monta uma grade de "etiquetas" (fundo verde-claro) com os nomes dos
    municípios vinculados — mais fácil de ler que uma linha corrida de texto."""
    estilo_chip = ParagraphStyle(
        "FichaSindChip", fontName="Helvetica", fontSize=10, leading=13,
        textColor=colors.HexColor("#23452E"),
    )
    nomes = [m.nome for m in municipios] or ["Nenhum município vinculado no momento"]
    linhas = [nomes[i:i + por_linha] for i in range(0, len(nomes), por_linha)]
    # completa a última linha pra manter o número de colunas constante
    if linhas and len(linhas[-1]) < por_linha:
        linhas[-1] += [""] * (por_linha - len(linhas[-1]))

    dados = [[Paragraph(nome, estilo_chip) if nome else "" for nome in linha] for linha in linhas]
    largura_col = (largura_total_mm / por_linha) * mm
    tabela = Table(dados, colWidths=[largura_col] * por_linha)
    estilo_cmds = [
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]
    for linha_idx, linha in enumerate(dados):
        for col_idx, celula in enumerate(linha):
            if celula != "":
                estilo_cmds.append(("BACKGROUND", (col_idx, linha_idx), (col_idx, linha_idx), colors.HexColor("#E7F0E8")))
    tabela.setStyle(TableStyle(estilo_cmds))
    return tabela


def gerar_ficha_sindicato_pdf(db, sindicato):
    """Gera uma ficha em PDF de UM sindicato, em retrato: dados atuais
    (presidente, período de gestão — com aviso se estiver vencendo — vice-
    presidente, contatos, endereço, foto) e municípios vinculados hoje em
    formato de etiquetas. Retorna bytes prontos para download."""

    municipios = municipios_do_sindicato(db, sindicato.id)
    referencia = melhor_referencia(municipios)

    verde_escuro = colors.HexColor("#23452E")
    cinza_texto = colors.HexColor("#6E6555")
    cor_cabecalho = _FICHA_VERDE if sindicato.ativo else colors.HexColor("#C0392B")

    titulo_style = ParagraphStyle(
        "FichaSindTitulo", fontName="Helvetica-Bold", fontSize=18,
        alignment=TA_CENTER, leading=21, textColor=colors.white,
    )
    subtitulo_style = ParagraphStyle(
        "FichaSindSubtitulo", fontName="Helvetica", fontSize=11,
        alignment=TA_CENTER, textColor=colors.white,
    )
    secao_titulo_style = ParagraphStyle(
        "FichaSindSecao", fontName="Helvetica-Bold", fontSize=13,
        textColor=verde_escuro, spaceBefore=2, spaceAfter=6,
    )
    rotulo_style = ParagraphStyle("FichaSindRot", fontName="Helvetica-Bold", fontSize=10.5, leading=14.5, textColor=cinza_texto)
    valor_style = ParagraphStyle("FichaSindVal", fontName="Helvetica-Bold", fontSize=12.5, leading=16, textColor=colors.black)
    valor_normal_style = ParagraphStyle("FichaSindValN", fontName="Helvetica", fontSize=11.5, leading=15, textColor=colors.black)

    largura_pagina_mm = 180  # A4 retrato, com margens de 15mm de cada lado

    story = []

    # ---------------- Cabeçalho ----------------
    nome_upper = sindicato.nome.upper()
    if "SINDICATO" in nome_upper:
        titulo_txt = nome_upper
    else:
        titulo_txt = f"SINDICATO RURAL DE {nome_upper}"
    subtitulo_txt = (referencia.regiao_faec.upper() if referencia and referencia.regiao_faec else "").strip()
    cabecalho_linhas = [[Paragraph(titulo_txt, titulo_style)]]
    if subtitulo_txt:
        cabecalho_linhas.append([Paragraph(subtitulo_txt, subtitulo_style)])
    cabecalho = Table(cabecalho_linhas, colWidths=[largura_pagina_mm * mm])
    estilo_cabecalho = [
        ("BACKGROUND", (0, 0), (-1, -1), cor_cabecalho),
        ("TOPPADDING", (0, 0), (-1, 0), 14),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 2 if subtitulo_txt else 14),
    ]
    if subtitulo_txt:
        estilo_cabecalho += [("TOPPADDING", (0, 1), (-1, 1), 0), ("BOTTOMPADDING", (0, 1), (-1, 1), 14)]
    cabecalho.setStyle(TableStyle(estilo_cabecalho))
    story.append(cabecalho)
    story.append(Spacer(1, 7 * mm))

    # ---------------- Presidente / gestão / foto ----------------
    linhas_dados = []
    cor_situacao_hex = "#23452E" if sindicato.ativo else "#C0392B"
    texto_situacao = "Ativo" if sindicato.ativo else "Inativo"
    linhas_dados.append(("Situação", f'<font color="{cor_situacao_hex}"><b>{texto_situacao}</b></font>' if not sindicato.ativo else texto_situacao, False))
    if referencia and referencia.presidente_nome:
        linhas_dados.append(("Presidente", _titulo_capitalizado(referencia.presidente_nome), True))
    linhas_dados.append(("Secretário(a)", _titulo_capitalizado(referencia.secretario_nome) if referencia and referencia.secretario_nome else "", False))

    inicio_txt = sindicato.gestao_inicio.strftime("%d/%m/%Y") if sindicato.gestao_inicio else "não informado"
    fim_txt = sindicato.gestao_fim.strftime("%d/%m/%Y") if sindicato.gestao_fim else "atual"
    periodo = f"Início: {inicio_txt} &nbsp;&nbsp;|&nbsp;&nbsp; Fim: {fim_txt}"
    aviso_txto, aviso_cor = _texto_situacao_mandato(sindicato)
    if aviso_txto:
        periodo_html = f'{periodo} &nbsp;<font color="{aviso_cor}"><b>{aviso_txto}</b></font>'
    else:
        periodo_html = periodo
    linhas_dados.append(("Mandato do presidente (início e fim)", periodo_html, False))

    # --- Contato do sindicato ---
    if referencia and (referencia.telefone1 or referencia.telefone2):
        linhas_dados.append(("Telefone(s) do sindicato", " / ".join(filter(None, [referencia.telefone1, referencia.telefone2])), False))
    if referencia and any([referencia.email1, referencia.email2, referencia.email3, referencia.email4]):
        linhas_dados.append(("E-mail(s) do sindicato", " / ".join(filter(None, [
            referencia.email1, referencia.email2, referencia.email3, referencia.email4])), False))

    # --- Vice-Presidente Regional (pessoa distinta, contato próprio) ---
    if referencia and referencia.vice_presidente_texto:
        linhas_dados.append(("Vice-Presidente Regional", _titulo_capitalizado(referencia.vice_presidente_texto), False))
    if referencia and referencia.telefone_vice_presidente:
        linhas_dados.append(("Telefone do Vice-Presidente Regional", referencia.telefone_vice_presidente, False))
    if referencia and referencia.email_vice_presidente:
        linhas_dados.append(("E-mail do Vice-Presidente Regional", referencia.email_vice_presidente, False))

    linhas_dados.append(("Endereço", _titulo_capitalizado(sindicato.endereco) or "não informado", False))
    if sindicato.latitude and sindicato.longitude:
        linhas_dados.append(("Localização", f"{sindicato.latitude}, {sindicato.longitude}", False))

    linhas_tabela_dados = []
    for rot, val, destaque in linhas_dados:
        estilo_valor = valor_style if destaque else valor_normal_style
        linhas_tabela_dados.append([Paragraph(rot, rotulo_style), Paragraph(val, estilo_valor)])

    estilo_info = TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("LINEBELOW", (0, 0), (-1, -2), 0.5, colors.HexColor("#E4DCC6")),
    ])

    foto_bytes = None
    if sindicato.foto_presidente_url:
        try:
            foto_bytes = supabase_storage.baixar_bytes(sindicato.foto_presidente_url)
        except Exception as erro:  # noqa: BLE001 — não vale travar o PDF por causa da foto
            print(f"[sindicatos_admin] aviso: não consegui baixar foto do Storage: {erro}")
    elif sindicato.foto_presidente:
        foto_bytes = sindicato.foto_presidente

    if foto_bytes:
        imagem = _imagem_pdf_da_foto(foto_bytes)
        largura_info = largura_pagina_mm - 40
        tabela_info = Table(linhas_tabela_dados, colWidths=[45 * mm, (largura_info - 45) * mm])
        tabela_info.setStyle(estilo_info)
        # Foto à direita, tabela de dados à esquerda
        cabecalho_com_foto = Table([[tabela_info, imagem]], colWidths=[largura_info * mm, 38 * mm])
        cabecalho_com_foto.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(cabecalho_com_foto)
    else:
        tabela_info = Table(linhas_tabela_dados, colWidths=[45 * mm, (largura_pagina_mm - 45) * mm])
        tabela_info.setStyle(estilo_info)
        story.append(tabela_info)

    story.append(Spacer(1, 8 * mm))

    # ---------------- Municípios vinculados ----------------
    story.append(Paragraph(f"Municípios vinculados ({len(municipios)})", secao_titulo_style))
    story.append(_chips_municipios(municipios, por_linha=1, largura_total_mm=largura_pagina_mm))

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        topMargin=15 * mm, bottomMargin=15 * mm,
        leftMargin=15 * mm, rightMargin=15 * mm,
        title=f"Ficha do Sindicato {sindicato.nome}",
    )
    doc.build(story)
    buffer.seek(0)
    return buffer.read()
