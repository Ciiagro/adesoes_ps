"""
Importa a planilha "Municipios_Ceara_Sindicatos_Unificado.xlsx" para o banco
(schema `sindicatos`, criado por `sindicatos_schema.sql`).

Rode isso UMA VEZ, depois de já ter executado o `sindicatos_schema.sql` no
SQL Editor do Supabase. É seguro rodar de novo (faz "upsert": atualiza quem
já existe, cria quem não existe) — mas rodar de novo por cima de dados já
editados manualmente no painel admin vai SOBRESCREVER essas edições, então
normalmente isso roda uma única vez, na carga inicial.

Uso:
    python importar_planilha_sindicatos.py
    python importar_planilha_sindicatos.py caminho/para/outra_planilha.xlsx

Não mexe no .env — usa a mesma DATABASE_URL que o resto do sistema já usa.
"""
import sys

import openpyxl

from database import SessionLocal, init_db
from models import SindicatoRural, MunicipioSindicato

ARQUIVO_PADRAO = "Municipios_Ceara_Sindicatos_Unificado.xlsx"
ABA = "Municipios_Sindicatos"

# Ordem das colunas na planilha original
COL_COD_IBGE = 0
COL_MUNICIPIO = 1
COL_REGIAO_FAEC = 2
COL_REGIAO_SEBRAE = 3
COL_REGIAO_PLANEJAMENTO = 4
COL_SINDICATO = 5
COL_VICE_PRESIDENTE = 6
COL_PRESIDENTE = 7
COL_TELEFONE1 = 8
COL_TELEFONE2 = 9
COL_EMAIL1 = 10
COL_EMAIL2 = 11
COL_EMAIL3 = 12
COL_EMAIL4 = 13


def limpar(valor):
    if valor is None:
        return None
    valor = str(valor).strip()
    return valor or None


def importar(caminho_arquivo):
    wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    ws = wb[ABA] if ABA in wb.sheetnames else wb.worksheets[0]

    db = SessionLocal()
    sindicatos_cache = {s.nome: s for s in db.query(SindicatoRural).all()}

    criados_sindicatos = 0
    criados_municipios = 0
    atualizados_municipios = 0

    for linha in ws.iter_rows(min_row=2, values_only=True):
        if linha[COL_COD_IBGE] is None:
            continue  # linha em branco no fim da planilha

        cod_ibge = int(linha[COL_COD_IBGE])
        nome_sindicato = limpar(linha[COL_SINDICATO])

        sindicato = None
        if nome_sindicato:
            sindicato = sindicatos_cache.get(nome_sindicato)
            if sindicato is None:
                sindicato = SindicatoRural(nome=nome_sindicato)
                db.add(sindicato)
                db.flush()  # garante o id antes de referenciar
                sindicatos_cache[nome_sindicato] = sindicato
                criados_sindicatos += 1

        municipio = db.get(MunicipioSindicato, cod_ibge)
        novo = municipio is None
        if novo:
            municipio = MunicipioSindicato(cod_ibge=cod_ibge)
            db.add(municipio)

        municipio.nome = limpar(linha[COL_MUNICIPIO])
        municipio.regiao_faec = limpar(linha[COL_REGIAO_FAEC])
        municipio.regiao_sebrae = limpar(linha[COL_REGIAO_SEBRAE])
        municipio.regiao_planejamento = limpar(linha[COL_REGIAO_PLANEJAMENTO])
        municipio.sindicato_id = sindicato.id if sindicato else None
        municipio.vice_presidente_texto = limpar(linha[COL_VICE_PRESIDENTE])
        municipio.presidente_nome = limpar(linha[COL_PRESIDENTE])
        municipio.telefone1 = limpar(linha[COL_TELEFONE1])
        municipio.telefone2 = limpar(linha[COL_TELEFONE2])
        municipio.email1 = limpar(linha[COL_EMAIL1])
        municipio.email2 = limpar(linha[COL_EMAIL2])
        municipio.email3 = limpar(linha[COL_EMAIL3])
        municipio.email4 = limpar(linha[COL_EMAIL4])
        municipio.atualizado_por = "importação inicial da planilha"

        if novo:
            criados_municipios += 1
        else:
            atualizados_municipios += 1

    db.commit()

    print(f"Sindicatos novos criados:      {criados_sindicatos}")
    print(f"Municípios novos criados:      {criados_municipios}")
    print(f"Municípios atualizados:        {atualizados_municipios}")
    print("Importação concluída.")


if __name__ == "__main__":
    init_db()  # garante que os models já registrados existam (create_all é idempotente)
    caminho = sys.argv[1] if len(sys.argv) > 1 else ARQUIVO_PADRAO
    importar(caminho)
